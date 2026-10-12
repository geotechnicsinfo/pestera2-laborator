# -*- coding: utf-8 -*-
"""
build_data.py — Laborator BBCGEOTECHNIC · Eoliene Peștera 2
Citește centralizatorul (cel mai recent „Centralizator*.xlsx” din folderul laboratorului)
și curbele din fișele Excel (foaia „Plot”) și scrie data.js pentru site.

Rulare (Windows):  py build_data.py
Scriptul stă în subfolderul site_pestera2 al folderului „pestera2 laborator”.
"""
import os, re, json, glob, datetime, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
LAB = os.path.dirname(HERE)
FISE = os.path.join(LAB, "fise excel")
OUT = os.path.join(HERE, "data.js")

# ---------- 1. centralizator ----------
cands = sorted(glob.glob(os.path.join(LAB, "Centralizator*.xlsx")), key=os.path.getmtime)
cands = [c for c in cands if not os.path.basename(c).startswith("~$")]
if not cands:
    sys.exit("Nu am găsit Centralizator*.xlsx în " + LAB)
CENTRAL = cands[-1]
print("Centralizator:", os.path.basename(CENTRAL))

# mapare coloane după textul din antet (rândul cu „Nr. raport”)
HEAD = [
    ("rep", "Nr. raport"), ("lot", "Lot"), ("bh", "Foraj"), ("prb", "Probă"), ("z", "Adâncime"),
    ("date", "Data raport"), ("t05", "Tip pământ SR EN ISO 14688-2:2005"),
    ("t18", "Tip pământ SR EN ISO 14688-2:2018"), ("obs", "Observații"), ("uscs", "Clasificare USCS"),
    ("A", "Argilă"), ("P", "Praf"), ("N", "Nisip"), ("Pi", "Pietriș"),
    ("w", "w"), ("wL", "wL"), ("wP", "wP"), ("IP", "IP"), ("IC", "IC"), ("cons", "Stare de"),
    ("IA", "IA"), ("g", "γ"), ("gd", "γd"), ("gsat", "γsat"), ("gs", "γs *"), ("n", "n"), ("e", "e"),
    ("Sr", "Sr"), ("M23n", "M2-3 nat"), ("M23i", "M2-3 inundat"), ("im3", "im3"), ("pu", "pu"),
    ("compr", "Clasificare compresibilitate"), ("tip", "Tip încercare"), ("v", "v"),
    ("phi", "φ"), ("c", "c"), ("tern", "Clasificare după diagrama"), ("pdf", "PDF în"),
]

wb = openpyxl.load_workbook(CENTRAL, data_only=True)
ws = wb["General"]
rows = list(ws.iter_rows(values_only=True))
hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "Nr. raport")
hdr = [str(h) if h is not None else "" for h in rows[hdr_i]]
full = [" ".join(h.split()) for h in hdr]          # antet complet, pe un rând
first = [h.split("\n")[0].strip() for h in hdr]     # primul rând al antetului
colmap = {}
for key, label in HEAD:
    j = next((j for j in range(len(hdr)) if full[j] == label), None)
    if j is None:
        j = next((j for j in range(len(hdr)) if first[j] == label and j not in colmap.values()), None)
    if j is not None:
        colmap[key] = j
missing = [k for k, _ in HEAD if k not in colmap]
if missing:
    print("ATENȚIE, coloane negăsite:", missing)

def clean(v):
    if v is None: return None
    if isinstance(v, str):
        s = v.strip()
        if s in ("", "-", "–", "#N/A", "N/A"): return None
        return s
    if isinstance(v, float):
        return round(v, 6)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%d.%m.%Y")
    return v

samples = []
for r in rows[hdr_i + 1:]:
    if not r or not isinstance(r[0], (int, float)):
        continue
    s = {k: clean(r[j]) for k, j in colmap.items()}
    s["rep"] = int(s["rep"])
    s["bh"] = str(s["bh"]).strip()
    samples.append(s)
print("Probe în centralizator:", len(samples))

# ---------- 2. fișe Excel: curbe ----------
pat = re.compile(r"^(C - )?(\d+) - BBC - .*?\.xlsx$", re.I)
files = {}
for f in os.listdir(FISE) if os.path.isdir(FISE) else []:
    m = pat.match(f)
    if not m or f.startswith("~$"): continue
    rep = int(m.group(2)); corr = bool(m.group(1))
    if rep not in files or corr:   # varianta „C - ” (corectată) are prioritate
        files[rep] = f

def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None

def read_curves(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Plot"]
    grid = {}
    for row in ws.iter_rows(min_row=1, max_row=300, max_col=17):
        for c in row:
            if c.value is not None:
                grid[(c.row, c.column)] = c.value
    wb.close()
    g = lambda r, c: grid.get((r, c))
    out = {}
    # granulometrie: antet „d (mm)” în col P (16)
    hr = next((r for r in range(56, 115) if str(g(r, 16) or "").startswith("d (mm)")), None)
    if hr:
        pts = []
        r = hr + 1
        while num(g(r, 16)) is not None and num(g(r, 17)) is not None:
            pts.append([g(r, 16), round(g(r, 17), 3)]); r += 1
        # elimină duplicatele de la capătul grosier (63 mm repetat)
        ded = []
        for p in pts:
            if not ded or abs(ded[-1][0] - p[0]) > 1e-9: ded.append(p)
        if len(ded) > 3: out["gran"] = ded
    # edometru: antet „Treapta de incarcare” în col A, între rândurile 173–235
    hr = next((r for r in range(173, 236) if str(g(r, 1) or "").startswith("Treapta de incarcare")), None)
    if hr:
        pts = []
        for r in range(hr + 1, hr + 26):
            if str(g(r, 1) or "").startswith("Intocmit"): break
            s = num(g(r, 1))
            if s is None: continue
            en, ei = num(g(r, 7)), num(g(r, 9))
            if en is None and ei is None: continue
            pts.append([s, None if en is None else round(en, 3), None if ei is None else round(ei, 3)])
        if pts: out["oed"] = pts
    # forfecare: „s=” în col A, rândurile 237–293
    hr = next((r for r in range(237, 294) if str(g(r, 1) or "").strip() == "s="), None)
    if hr:
        sig = [num(g(hr, c)) for c in (2, 4, 6)]
        tau = [num(g(hr + 1, c)) for c in (2, 4, 6)]
        pts = [[a, round(b, 2)] for a, b in zip(sig, tau) if a is not None and b is not None]
        if len(pts) >= 2:
            out["shear"] = pts
            out["shear_phi"] = num(g(hr - 5, 2)); out["shear_c"] = num(g(hr - 4, 2))
    return out

# ---------- 3. forfecare directă din PDF-urile buletinelor (sursa finală) ----------
# Rezultatele se păstrează în forfecare_pdf.json; pentru PDF-uri noi se încearcă
# pdftotext (Poppler) sau pypdf (pip install pypdf). Fără ele rămân valorile din cache/fișă.
import subprocess
CACHE = os.path.join(HERE, "forfecare_pdf.json")
cache = json.load(open(CACHE, encoding="utf-8")) if os.path.exists(CACHE) else {}
pdfs = {}
for p in glob.glob(os.path.join(LAB, "*", "*.pdf")):
    m = re.match(r"^(C - )?(\d+) - BBC", os.path.basename(p))
    if m and (int(m.group(2)) not in pdfs or m.group(1)):
        pdfs[int(m.group(2))] = p

def pdf_text(p):
    try:
        return subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True,
                              encoding="utf-8", errors="ignore", timeout=60).stdout
    except Exception:
        pass
    try:
        import pypdf
        return "\n".join(pg.extract_text() or "" for pg in pypdf.PdfReader(p).pages)
    except Exception:
        return ""

def shear_from_pdf(p):
    t = pdf_text(p)
    i = t.find("forfecare directa /")
    i = t.find("Raport de incercare - Incercarea de forfecare", max(i, 0))
    if i < 0: return None
    t = t[i:i + 4000]
    fl = r"([\d]+[.,]?\d*)"
    ms = re.search(r"s=\s*" + fl + r"\s+" + fl + r"\s+" + fl, t)
    mt = re.search(r"tf=\s*" + fl + r"\s+" + fl + r"\s+" + fl, t)
    if not (ms and mt):   # pypdf pierde simbolurile σ/τ: se caută rândurile „… m (g)” și „… md (g)”
        ms = re.search(fl + r"\s+" + fl + r"\s+" + fl + r"\s+m \(g\)", t)
        mt = re.search(fl + r"\s+" + fl + r"\s+" + fl + r"\s+md \(g\)", t)
    if not (ms and mt): return None
    f = lambda x: float(x.replace(",", "."))
    out = {"shear": [[f(ms.group(k)), f(mt.group(k))] for k in (1, 2, 3)]}
    mp = re.search(r"f=\s*" + fl, t); mc = re.search(r"c=\s*" + fl, t)
    if mp: out["phi"] = f(mp.group(1))
    if mc: out["c"] = f(mc.group(1))
    mtyp = re.search(r"\b(CU|CD|UU)\s+(natural|saturat)", t)
    if mtyp: out["tip"] = mtyp.group(1) + " " + mtyp.group(2)
    return out

n_pdf = 0
for s in samples:
    if s.get("phi") is None: continue
    key = str(s["rep"])
    if key not in cache and s["rep"] in pdfs:
        r = shear_from_pdf(pdfs[s["rep"]])
        if r: cache[key] = r
    if key in cache:
        s["shear_pdf"] = cache[key]; n_pdf += 1
json.dump(cache, open(CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("Forfecare din PDF:", n_pdf, "/", sum(1 for s in samples if s.get("phi") is not None))

n_ok = 0
for s in samples:
    f = files.get(s["rep"])
    s["fisa"] = f
    if not f: continue
    try:
        cv = read_curves(os.path.join(FISE, f))
        s.update(cv); n_ok += 1
    except Exception as ex:
        print("Eroare la", f, ex)
print("Fișe citite:", n_ok, "/", len(samples))

for s in samples:   # PDF-ul are prioritate față de foaia Plot pentru punctele τ–σ
    sp = s.pop("shear_pdf", None)
    s.pop("shear_phi", None); s.pop("shear_c", None)
    if sp:
        s["shear"] = sp["shear"]
        if sp.get("tip"): s["tip_pdf"] = sp["tip"]

data = {
    "meta": {
        "proiect": "Eoliene Peștera 2",
        "beneficiar": "GEOSOND SA",
        "amplasament": "localitatea Peștera, jud. Constanța",
        "laborator": "BBCGEOTECHNIC",
        "sursa": os.path.basename(CENTRAL),
        "generat": datetime.datetime.now().strftime("%d.%m.%Y %H:%M"),
        "nfise": n_ok,
    },
    "samples": samples,
}
with open(OUT, "w", encoding="utf-8") as fo:
    fo.write("window.LAB = ")
    json.dump(data, fo, ensure_ascii=False, separators=(",", ":"))
    fo.write(";\n")
print("Scris:", OUT, os.path.getsize(OUT) // 1024, "KB")
